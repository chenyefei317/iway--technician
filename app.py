import datetime
import io
import os
import re
import zipfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor, Pt
import pandas as pd
import qrcode
from PIL import Image
import requests
import streamlit as st
from streamlit_drawable_canvas import st_canvas

# ================= 1. 页面配置与初始化 =================
st.set_page_config(
    page_title="员工职业危害告知书和转岗培训记录表签收平台",
    layout="centered",
    initial_sidebar_state="expanded",
)

# 注入华文宋体全局样式与隐藏默认元素
hide_streamlit_style = """
    <style>
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    html, body, [class*="css"] {
        font-family: "华文宋体", SimSun, serif;
    }
    </style>
    """
st.markdown(hide_streamlit_style, unsafe_allow_html=True)


# 智能模糊查找 Word 文件函数（通用）
def find_docx_file(folder, keyword1, keyword2=None):
  if not os.path.exists(folder):
    return None
  for filename in os.listdir(folder):
    if keyword2:
      if (
          keyword1 in filename
          and keyword2 in filename
          and filename.lower().endswith(".docx")
      ):
        return os.path.join(folder, filename)
    else:
      if keyword1 in filename and filename.lower().endswith(".docx"):
        return os.path.join(folder, filename)
  return None


# ================= 2. 百度网盘自动上传函数（带 OAuth2 自动刷新） =================
def refresh_baidu_access_token():
  try:
    client_id = st.secrets.get("BAIDU_CLIENT_ID", "")
    client_secret = st.secrets.get("BAIDU_CLIENT_SECRET", "")
    refresh_token = st.secrets.get("BAIDU_REFRESH_TOKEN", "")
    if not client_id or not client_secret or not refresh_token:
      return None
    token_url = "https://pan.baidu.com/oauth/2.0/token"
    params = {
        "grant_type": "refresh_token",
        "refresh_token": refresh_token,
        "client_id": client_id,
        "client_secret": client_secret,
    }
    response = requests.get(token_url, params=params)
    res_data = response.json()
    return res_data.get("access_token")
  except Exception:
    return None


def upload_to_baidu_netdisk_with_auto_refresh(file_bytes, remote_filename):
  access_token = st.secrets.get("BAIDU_ACCESS_TOKEN", "")
  if not access_token:
    return (
        False,
        "未配置网盘凭证，文件已在本地生成并可通过网页下载/ZIP打包保存。",
    )

  sub_folder = "EHS签字档案"
  target_path = f"/apps/慧瑞EHS合规档案/{sub_folder}/{remote_filename}"

  def send_upload_request(token):
    upload_url = f"https://pan.baidu.com/rest/2.0/xpan/file?method=upload&access_token={token}&path={target_path}&uploadid=&file=1"
    files = {"file": (remote_filename, file_bytes)}
    return requests.post(upload_url, files=files).json()

  result = send_upload_request(access_token)
  if "errno" in result and result["errno"] in [110, 111]:
    new_token = refresh_baidu_access_token()
    if new_token:
      result = send_upload_request(new_token)
    else:
      return False, "Token 已过期且自动刷新失败。"

  if "errno" in result and result["errno"] == 0:
    return True, f"成功同步至网盘：/apps/慧瑞EHS合规档案/{sub_folder}/"
  else:
    return False, f"网盘上传失败: {result.get('error_msg', '未知错误')}"


# ================= 3. 侧边栏：Logo、微信分享与模板下载 =================
with st.sidebar:
  try:
    st.image("logo.png", width=160)
  except Exception:
    st.image(
        "https://upload.wikimedia.org/wikipedia/commons/thumb/c/c5/Ikea_logo.svg/800px-Ikea_logo.svg.png",
        width=160,
    )

  st.markdown("### 📱 微信扫码与分享")
  st.write("已自动关联您的云端网址，二维码将实时更新供手机扫码填报。")

  app_url = st.text_input(
      "应用公网链接 (URL)", value="https://iway--technician.streamlit.app"
  )

  if app_url:
    qr = qrcode.make(app_url)
    img_buffer = io.BytesIO()
    qr.save(img_buffer, format="PNG")
    st.image(
        Image.open(img_buffer), caption="微信扫码快速填报通道", width=160
    )
    st.info(
        "💡 **提示**：将上方链接复制并发送至微信工作群，员工即可手机端签收。"
    )

  st.markdown("---")
  st.markdown("### 📥 常用制度模板快捷下载")

  # 培训记录表模板下载
  training_path_dl = find_docx_file(
      "员工转岗安全与职业健康培训记录表", "转岗安全与职业健康培训记录表"
  )
  if training_path_dl and os.path.exists(training_path_dl):
    with open(training_path_dl, "rb") as ft:
      st.download_button(
          label="📄 转岗培训记录表模板",
          data=ft.read(),
          file_name="员工转岗安全与职业健康培训记录表.docx",
      )

  # 侧边栏职业危害告知书 Word 下载
  st.write("**职业危害告知书 Word 下载：**")
  side_company = st.selectbox(
      "选择公司：",
      ["安徽恒林", "大连宜家", "东莞时兴", "福建龙竹", "福建双翼"],
      key="side_comp",
  )
  side_stage = st.selectbox(
      "选择阶段：", ["上岗前", "在岗期间"], key="side_stage"
  )

  side_hazard_path = find_docx_file(
      "职业危害告知书", side_company, side_stage
  )
  if side_hazard_path and os.path.exists(side_hazard_path):
    with open(side_hazard_path, "rb") as fsh:
      st.download_button(
          label=f"📥 下载选中的告知书",
          data=fsh.read(),
          file_name=f"职业危害告知书 - {side_company}（{side_stage}）.docx",
          key="side_dl_hazard_btn",
      )
  else:
    st.warning("⚠️ 暂未找到对应的 Word 模板")

# ================= 4. 主界面逻辑（Logo在左侧，主标题单独一行） =================
col_logo, col_title = st.columns([1, 6])
with col_logo:
  try:
    st.image("logo.png", width=110)
  except Exception:
    st.image(
        "https://upload.wikimedia.org/wikipedia/commons/thumb/c/c5/Ikea_logo.svg/800px-Ikea_logo.svg.png",
        width=110,
    )
with col_title:
  st.markdown("## 员工职业危害告知书和转岗培训记录表签收平台")

st.markdown(
    "请仔细阅读下方各项内容，勾选确认并在底部完成手写签收与手写日期。系统将自动把您的亲笔签名与手写日期嵌入对应的"
    " Word 正式档案中。"
)

# 基础信息录入
st.subheader("1. 员工基本信息")
col1, col2 = st.columns(2)
with col1:
  emp_name = st.text_input("员工姓名 (必填)：")
with col2:
  emp_id = st.text_input(
      "身份证号 (必填，须满18位)：",
      help="请输入标准的 18 位中国居民身份证号码",
  )

st.write("---")
st.markdown("### 📂 待签收项目清单")

# --- 项目一：职业危害告知书（Word版本，非必选项） ---
st.subheader("⚠️ 项目一：职业危害告知书 (可选)")

col_c, col_s = st.columns(2)
with col_c:
  company_choice = st.selectbox(
      "选择公司/站点：",
      ["安徽恒林", "大连宜家", "东莞时兴", "福建龙竹", "福建双翼"],
  )
with col_s:
  stage_choice = st.selectbox("选择告知阶段：", ["上岗前", "在岗期间"])

hazard_version = f"职业危害告知书 - {company_choice}（{stage_choice}）"
hazard_folder = "职业危害告知书"
hazard_path = find_docx_file(hazard_folder, company_choice, stage_choice)

try:
  if hazard_path and os.path.exists(hazard_path):
    with open(hazard_path, "rb") as f:
      hazard_docx_data = f.read()
  else:
    raise FileNotFoundError
except FileNotFoundError:
  doc_temp = Document()
  doc_temp.add_heading(hazard_version, level=1)
  doc_temp.add_paragraph(
      f"【系统提示】在 '{hazard_folder}' 文件夹中未找到匹配的 '.docx'"
      " 文件，请确认已上传至 GitHub。"
  )
  temp_io = io.BytesIO()
  doc_temp.save(temp_io)
  hazard_docx_data = temp_io.getvalue()
  hazard_path = None

st.info(f"您当前查阅的是：【{hazard_version}】。")
st.download_button(
    label=f"📥 下载《{hazard_version}.docx》",
    data=hazard_docx_data,
    file_name=f"{hazard_version}.docx",
    mime=(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
)
c_hazard = st.checkbox(
    f"【可选确认】本人已阅读并充分了解《{hazard_version}》的相关职业危害与防护要求，承诺在工作中严格落实。"
)

with st.expander("📚 附加查阅：全套《职业危害告知书》Word 快捷下载专区"):
  st.write(
      "如需查阅或下载其他公司/阶段的职业危害告知书，可直接点击下方按钮："
  )
  if os.path.exists(hazard_folder):
    all_hazard_files = [
        f for f in os.listdir(hazard_folder) if f.lower().endswith(".docx")
    ]
    for hf in all_hazard_files:
      full_hf_path = os.path.join(hazard_folder, hf)
      with open(full_hf_path, "rb") as fh:
        st.download_button(
            label=f"📥 下载：{hf}",
            data=fh.read(),
            file_name=hf,
            key=f"dl_all_docx_{hf}",
        )
  else:
    st.write("暂无其他告知书 Word 文件。")

st.write("---")

# --- 项目二：员工转岗安全与职业健康培训记录表 ---
st.subheader("🎓 项目二：员工转岗安全与职业健康培训记录表")
training_folder = "员工转岗安全与职业健康培训记录表"
training_path = find_docx_file(training_folder, "转岗安全与职业健康培训记录表")

try:
  if training_path and os.path.exists(training_path):
    with open(training_path, "rb") as f:
      training_docx_data = f.read()
  else:
    raise FileNotFoundError
except FileNotFoundError:
  doc_temp2 = Document()
  doc_temp2.add_heading("员工转岗安全与职业健康培训记录表", level=1)
  doc_temp2.add_paragraph(
      "【系统提示】未在 '员工转岗安全与职业健康培训记录表' 文件夹中找到对应的"
      " '.docx' 文件。"
  )
  temp_io2 = io.BytesIO()
  doc_temp2.save(temp_io2)
  training_docx_data = temp_io2.getvalue()
  training_path = None

st.markdown("请查阅以下标准培训记录表文件：")
st.download_button(
    label="📥 下载《员工转岗安全与职业健康培训记录表.docx》",
    data=training_docx_data,
    file_name="员工转岗安全与职业健康培训记录表.docx",
    mime=(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
)
c_training = st.checkbox(
    "【须确认】本人已完成《员工转岗安全与职业健康培训记录表》所含全部课程的学习，熟知岗位危险源与操作规程。"
)

# ================= 5. 手写签名与手写日期栏（并排双画布，动态关联当前系统日期） =================
current_date_str = datetime.date.today().strftime("%Y年%m月%d日")

st.write("---")
st.subheader("✍️ 3. 员工手写签名与手写日期栏")
st.markdown(
    f"**请在左侧手写签名，并在右侧手写日期（注：当前系统日期为"
    f" {current_date_str}，请按此手写日期）：**"
)

col_sig, col_date = st.columns(2)
with col_sig:
  st.markdown("**手写签名：**")
  canvas_result = st_canvas(
      stroke_width=4,
      stroke_color="#000000",
      background_color="#F8F9FA",
      height=200,
      width=320,
      drawing_mode="freedraw",
      key="canvas_sig",
      return_image_data=True,
  )
with col_date:
  st.markdown("**手写日期栏（请手写当前日期）：**")
  canvas_date_result = st_canvas(
      stroke_width=3,
      stroke_color="#000000",
      background_color="#F8F9FA",
      height=200,
      width=320,
      drawing_mode="freedraw",
      key="canvas_date",
      return_image_data=True,
  )


# ================= 6. 辅助函数：向 Word 模板文末追加签名与日期 =================
def append_signature_to_docx(
    template_path, default_title, sig_image_io, date_image_io
):
  if template_path and os.path.exists(template_path):
    try:
      doc = Document(template_path)
    except Exception:
      doc = Document()
      doc.add_heading(default_title, level=1)
      doc.add_paragraph("（提示：模板文件读取异常，此为生成的标准确认单）")
  else:
    doc = Document()
    doc.add_heading(default_title, level=1)
    doc.add_paragraph("（提示：未找到对应的 .docx 模板文件）")

  # 统一设置华文宋体
  for p in doc.paragraphs:
    for r in p.runs:
      r.font.name = "华文宋体"
      r.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")

  p_line = doc.add_paragraph("--------------------------------------------------")
  p_line.paragraph_format.space_before = Pt(2)
  p_line.paragraph_format.space_after = Pt(2)

  p_confirm = doc.add_paragraph()
  p_confirm.paragraph_format.space_before = Pt(0)
  p_confirm.paragraph_format.space_after = Pt(2)
  run_c = p_confirm.add_run(
      f"【员工签收确认】 姓名：{emp_name} | 身份证号：{emp_id}\n"
      f"本人已仔细阅读并充分理解上述内容，承诺在工作中严格遵守各项安全防范及操作规程。"
  )
  run_c.font.name = "华文宋体"
  run_c.font.size = Pt(10.5)
  run_c.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")

  # 使用表格将手写签名与手写日期并排显示在底部
  table = doc.add_table(rows=1, cols=2)
  table.autofit = False

  cell_sig = table.cell(0, 0)
  p1 = cell_sig.paragraphs[0]
  r1 = p1.add_run("员工手写签名：\n")
  r1.font.name = "华文宋体"
  r1.font.size = Pt(10)
  r1.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")
  p1.add_run().add_picture(sig_image_io, width=Inches(1.6))
  sig_image_io.seek(0)

  cell_date = table.cell(0, 1)
  p2 = cell_date.paragraphs[0]
  r2 = p2.add_run("手写日期：\n")
  r2.font.name = "华文宋体"
  r2.font.size = Pt(10)
  r2.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")
  p2.add_run().add_picture(date_image_io, width=Inches(1.6))
  date_image_io.seek(0)

  buffer = io.BytesIO()
  doc.save(buffer)
  buffer.seek(0)
  return buffer


# ================= 7. 提交校验与生成带签名的档案 =================
if st.button(
    "📁 确认无误，一键签收并生成带签名的合规档案", use_container_width=True
):
  is_canvas_empty = canvas_result.image_data is None or (
      canvas_result.json_data is not None
      and len(canvas_result.json_data.get("objects", [])) == 0
  )
  is_date_empty = canvas_date_result.image_data is None or (
      canvas_date_result.json_data is not None
      and len(canvas_date_result.json_data.get("objects", [])) == 0
  )

  id_pattern = re.compile(r"^\d{17}[\dXx]$")

  if not emp_name.strip() or not emp_id.strip():
    st.error("❌ 拦截：请完整填写【员工姓名】与【身份证号】！")
  elif not id_pattern.match(emp_id.strip()):
    st.error(
        "❌ 拦截：身份证号必须为严格的 **18 位**数字（末尾可为大写 X）！"
    )
  elif not c_hazard and not c_training:
    st.error(
        "❌ 拦截：请至少勾选并完成一项签收（职业危害告知书或转岗培训记录表）！"
    )
  elif is_canvas_empty:
    st.warning("⚠️ 拦截：请在左侧画板完成手写签名后再提交。")
  elif is_date_empty:
    st.warning("⚠️ 拦截：请在右侧手写日期栏内完成手写日期后再提交！")
  else:
    st.success(
        "✅ 签收成功！系统已成功生成您的专属带签名 Word 合规档案。"
    )

    signature_img = Image.fromarray(
        canvas_result.image_data.astype("uint8"), "RGBA"
    )
    sig_io = io.BytesIO()
    signature_img.save(sig_io, format="PNG")
    sig_io.seek(0)

    date_img = Image.fromarray(
        canvas_date_result.image_data.astype("uint8"), "RGBA"
    )
    date_io = io.BytesIO()
    date_img.save(date_io, format="PNG")
    date_io.seek(0)

    # 动态生成用户勾选的文件并打包
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
      if c_hazard and hazard_path and os.path.exists(hazard_path):
        signed_hazard_buffer = append_signature_to_docx(
            hazard_path, f"{hazard_version} 签收单", sig_io, date_io
        )
        hazard_filename_cloud = (
            f"{hazard_version}_{emp_name}_{emp_id[-4:]}_已签字.docx"
        )
        zip_file.writestr(
            hazard_filename_cloud, signed_hazard_buffer.getvalue()
        )
        upload_to_baidu_netdisk_with_auto_refresh(
            signed_hazard_buffer.getvalue(), hazard_filename_cloud
        )

      if c_training and training_path:
        signed_training_buffer = append_signature_to_docx(
            training_path,
            "员工转岗安全与职业健康培训记录表 签收单",
            sig_io,
            date_io,
        )
        training_filename_cloud = (
            f"员工转岗培训记录表_{emp_name}_{emp_id[-4:]}_已签字.docx"
        )
        zip_file.writestr(
            training_filename_cloud, signed_training_buffer.getvalue()
        )
        upload_to_baidu_netdisk_with_auto_refresh(
            signed_training_buffer.getvalue(), training_filename_cloud
        )

      # 保存签名及日期原图
      img_byte_arr = io.BytesIO()
      signature_img.save(img_byte_arr, format="PNG")
      zip_file.writestr(
          f"手写签名原图_{emp_name}.png", img_byte_arr.getvalue()
      )

      date_byte_arr = io.BytesIO()
      date_img.save(date_byte_arr, format="PNG")
      zip_file.writestr(f"手写日期原图_{emp_name}.png", date_byte_arr.getvalue())

    zip_buffer.seek(0)

    st.markdown("---")
    st.success(
        "🎉 您的专属带签名 Word 合规档案已打包完毕，点击下方按钮即可下载！"
    )

    col_d1, col_d2 = st.columns(2)
    if c_hazard and hazard_path and os.path.exists(hazard_path):
      with col_d1:
        st.download_button(
            label="📄 下载带签名的告知书 (.docx)",
            data=signed_hazard_buffer.getvalue(),
            file_name=f"{hazard_version}_{emp_name}_已签字.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
        )
    if c_training and training_path:
      with col_d2:
        st.download_button(
            label="📄 下载带签名的培训表 (.docx)",
            data=signed_training_buffer.getvalue(),
            file_name=f"员工转岗培训记录表_{emp_name}_已签字.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
        )

    st.markdown("---")
    st.download_button(
        label="📥 一键打包下载全部签收档案 (.ZIP)",
        data=zip_buffer,
        file_name=f"安全合规档案_{emp_name}.zip",
        mime="application/zip",
        use_container_width=True,
    )

    st.balloons()

# ================= 8. 底部版权与开发者声明 =================
st.markdown("---")
st.markdown(
    "<div style='text-align: center; color: gray; font-size: 14px;'>"
    "内部使用，严禁商业用途 | 开发者：陈野菲 Yefei"
    "</div>",
    unsafe_allow_html=True,
)
